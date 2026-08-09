"""
Report Export Service (PDF / Excel)

Renders the JSON reports produced by ``report_service`` into downloadable
PDF (reportlab) and Excel (openpyxl) documents.

Public API returns raw bytes ready for a streaming ``Response``:
    anomaly_report_pdf_bytes / anomaly_report_xlsx_bytes
    employee_report_pdf_bytes / employee_report_xlsx_bytes
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

from sqlalchemy.orm import Session

from app.services.report_service import (
    generate_anomaly_report,
    generate_employee_report,
)

SYSTEM_NAME = "Insider Threat Behavioral Intelligence System (ITBIS)"


# ── Data helpers ────────────────────────────────────────────────


def _anomaly_data(db: Session, days: int) -> dict[str, Any]:
    return generate_anomaly_report(db, days=days, include_threat_assessment=True)


def _employee_data(db: Session, employee_id: str, days: int) -> dict[str, Any]:
    report = generate_employee_report(db, employee_id, days)
    if "error" in report:
        raise ValueError(report["error"])
    return report


# ── PDF (reportlab) ─────────────────────────────────────────────


def anomaly_report_pdf_bytes(db: Session, days: int = 30) -> bytes:
    """Render the organization anomaly report as a PDF byte stream."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    data = _anomaly_data(db, days)

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=A4, topMargin=0.7 * inch, bottomMargin=0.7 * inch
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleX", parent=styles["Title"], fontSize=16, leading=20,
        textColor=colors.HexColor("#0f172a"),
    )
    h2 = ParagraphStyle(
        "H2X", parent=styles["Heading2"], fontSize=12, leading=16,
        spaceBefore=10, spaceAfter=4, textColor=colors.HexColor("#0f172a"),
    )
    body = ParagraphStyle(
        "BodyX", parent=styles["BodyText"], fontSize=9, leading=12,
    )
    small = ParagraphStyle(
        "SmallX", parent=body, fontSize=8, textColor=colors.HexColor("#475569"),
    )

    story: list[Any] = []
    story.append(Paragraph(SYSTEM_NAME, title_style))
    story.append(
        Paragraph(
            f"Organization Anomaly Report — {days}-day lookback",
            styles["Normal"],
        )
    )
    story.append(
        Paragraph(
            f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')} "
            f"| Period: {data['report_period']['start'][:10]} to {data['report_period']['end'][:10]}",
            small,
        )
    )
    story.append(Spacer(1, 0.15 * inch))

    # ── Summary ─────────────────────────────────────────────
    summary = data["summary"]
    summary_rows = [
        ["Employees", str(summary["total_employees"])],
        ["Total Alerts", str(summary["total_alerts"])],
        ["Open Alerts", str(summary["open_alerts"])],
        ["Critical Alerts", str(summary["critical_alerts"])],
        ["Total Incidents", str(summary["total_incidents"])],
        ["Total Activity Logs", f"{summary['total_activity_logs']:,}"],
    ]
    story.append(Paragraph("Executive Summary", h2))
    story.append(_pdf_table(["Metric", "Value"], summary_rows))
    story.append(Spacer(1, 0.12 * inch))

    # ── Alert breakdown ─────────────────────────────────────
    story.append(Paragraph("Alert Analysis", h2))
    alerts = data["alerts"]
    story.append(
        Paragraph(
            f"Total: {alerts['total']} · Trend: {alerts['trend']}",
            body,
        )
    )
    severity_str = ", ".join(
        f"{k} ({v})" for k, v in alerts["by_severity"].items()
    )
    story.append(Paragraph(f"Severity breakdown: {severity_str or 'No alerts'}", body))
    top_types = ", ".join(
        f"{k.replace('_', ' ')} ({v})"
        for k, v in list(alerts["by_anomaly_type"].items())[:10]
    )
    story.append(Paragraph(f"Top anomaly types: {top_types or 'None'}", body))
    story.append(Spacer(1, 0.12 * inch))

    # ── Incidents ───────────────────────────────────────────
    story.append(Paragraph("Incident Analysis", h2))
    inc = data["incidents"]
    incident_rows = [
        ["Total Incidents", str(inc["total"])],
        ["Avg Resolution (hours)", str(inc.get("avg_resolution_hours") or "—")],
        ["Escalation Rate (%)", str(inc.get("escalation_rate", 0))],
    ]
    incident_rows += [[k, str(v)] for k, v in inc["by_status"].items()]
    story.append(_pdf_table(["Metric", "Value"], incident_rows))
    story.append(Spacer(1, 0.12 * inch))

    # ── Risk distribution ───────────────────────────────────
    story.append(Paragraph("Risk Score Distribution", h2))
    dist_rows = [[k.capitalize(), str(v)] for k, v in data["risk_distribution"].items()]
    story.append(_pdf_table(["Risk Level", "Employees"], dist_rows))
    story.append(Spacer(1, 0.12 * inch))

    # ── High risk employees ─────────────────────────────────
    story.append(Paragraph("High Risk Employees", h2))
    hr_rows = [
        [e["name"], e.get("department") or "—", e.get("designation") or "—",
         str(e["risk_score"]), e["risk_level"]]
        for e in data["high_risk_employees"]
    ]
    if hr_rows:
        story.append(_pdf_table(["Name", "Department", "Designation", "Risk", "Level"], hr_rows))
    else:
        story.append(Paragraph("No high-risk employees.", body))

    # ── Threat assessment ───────────────────────────────────
    ta = data.get("threat_assessment")
    if ta:
        story.append(Spacer(1, 0.12 * inch))
        story.append(Paragraph("Threat Assessment", h2))
        story.append(
            Paragraph(
                f"Employees assessed: {ta['total_assessed']} · "
                f"Average threat score: {ta['average_threat_score']}",
                body,
            )
        )
        top_rows = [
            [t.get("employee_name") or "—", t.get("department") or "—",
             str(t.get("threat_score", "")), t.get("threat_level") or "—"]
            for t in ta["top_threats"][:15]
        ]
        if top_rows:
            story.append(_pdf_table(["Name", "Department", "Score", "Level"], top_rows))

    doc.build(story)
    return buffer.getvalue()


def employee_report_pdf_bytes(
    db: Session, employee_id: str, days: int = 30
) -> bytes:
    """Render a single-employee report as a PDF byte stream."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table

    data = _employee_data(db, employee_id, days)
    emp = data["employee"]

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=0.7 * inch, bottomMargin=0.7 * inch)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleX", parent=styles["Title"], fontSize=15, leading=19,
        textColor=colors.HexColor("#0f172a"),
    )
    h2 = ParagraphStyle(
        "H2X", parent=styles["Heading2"], fontSize=11, leading=15,
        spaceBefore=10, spaceAfter=4, textColor=colors.HexColor("#0f172a"),
    )
    body = ParagraphStyle("BodyX", parent=styles["BodyText"], fontSize=9, leading=12)
    small = ParagraphStyle(
        "SmallX", parent=body, fontSize=8, textColor=colors.HexColor("#475569"),
    )

    story: list[Any] = []
    story.append(Paragraph(SYSTEM_NAME, title_style))
    story.append(Paragraph(f"Employee Report — {emp['name']}", styles["Normal"]))
    story.append(
        Paragraph(
            f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')} · "
            f"{emp.get('department') or '—'} · {emp.get('designation') or '—'}",
            small,
        )
    )
    story.append(Spacer(1, 0.15 * inch))

    summary_rows = [
        ["Total Alerts", str(data["total_alerts"])],
        ["Total Activity Logs", f"{data['total_activity_logs']:,}"],
    ]
    if data.get("latest_risk_score"):
        lr = data["latest_risk_score"]
        summary_rows.append(["Latest Risk Score", f"{lr['score']} ({lr['level']})"])
    story.append(Paragraph("Summary", h2))
    story.append(_pdf_table(["Metric", "Value"], summary_rows))
    story.append(Spacer(1, 0.12 * inch))

    if data.get("alert_severity_breakdown"):
        story.append(Paragraph("Alert Severity Breakdown", h2))
        sev_rows = [[k.capitalize(), str(v)] for k, v in data["alert_severity_breakdown"].items()]
        story.append(_pdf_table(["Severity", "Count"], sev_rows))
        story.append(Spacer(1, 0.12 * inch))

    if data.get("alert_type_breakdown"):
        story.append(Paragraph("Alert Type Breakdown", h2))
        type_rows = [
            [k.replace("_", " "), str(v)]
            for k, v in data["alert_type_breakdown"].items()
        ]
        story.append(_pdf_table(["Anomaly Type", "Count"], type_rows))
        story.append(Spacer(1, 0.12 * inch))

    ta = data.get("threat_assessment") or {}
    if ta:
        story.append(Paragraph("Threat Assessment", h2))
        story.append(
            Paragraph(
                f"Threat score: {ta.get('threat_score', '—')} · Level: {ta.get('threat_level', '—')}",
                body,
            )
        )
        model_rows = [
            [name.replace("_", " "), str(m.get("score", "")), m.get("level", "")]
            for name, m in (ta.get("model_scores") or {}).items()
        ]
        if model_rows:
            story.append(_pdf_table(["Model", "Score", "Level"], model_rows))

    doc.build(story)
    return buffer.getvalue()


def _pdf_table(headers: list[str], rows: list[list[str]]) -> Table:
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle

    t = Table([headers] + rows, repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, 0), 8),
                ("FONTSIZE", (0, 1), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return t


# ── Excel (openpyxl) ────────────────────────────────────────────


def anomaly_report_xlsx_bytes(db: Session, days: int = 30) -> bytes:
    """Render the organization anomaly report as an Excel workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    data = _anomaly_data(db, days)
    wb = Workbook()

    header_fill = PatternFill("solid", fgColor="1E293B")
    header_font = Font(color="FFFFFF", bold=True)

    def _style_header(ws, row: int, cols: int) -> None:
        for c in range(1, cols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = header_fill
            cell.font = header_font

    # ── Summary sheet ───────────────────────────────────────
    ws = wb.active
    ws.title = "Summary"
    ws.append(["ITBIS Organization Anomaly Report"])
    ws.append([f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"])
    ws.append([f"Lookback: {days} days"])
    ws.append([])
    ws.append(["Metric", "Value"])
    _style_header(ws, 5, 2)
    for k, v in data["summary"].items():
        ws.append([k.replace("_", " ").title(), v])
    ws.append([])
    ws.append(["Risk Distribution", ""])
    for level, count in data["risk_distribution"].items():
        ws.append([f"  {level.capitalize()}", count])
    for col in ("A", "B"):
        ws.column_dimensions[col].width = 30

    # ── Alerts sheet ────────────────────────────────────────
    ws2 = wb.create_sheet("Alerts")
    ws2.append(["Alert Analysis"])
    ws2.append([])
    ws2.append(["By Severity", "Count"])
    _style_header(ws2, 3, 2)
    for k, v in data["alerts"]["by_severity"].items():
        ws2.append([k, v])
    ws2.append([])
    ws2.append(["By Anomaly Type", "Count"])
    _style_header(ws2, 3 + len(data["alerts"]["by_severity"]) + 2, 2)
    for k, v in data["alerts"]["by_anomaly_type"].items():
        ws2.append([k, v])
    ws2.column_dimensions["A"].width = 28

    # ── Incidents sheet ─────────────────────────────────────
    ws3 = wb.create_sheet("Incidents")
    ws3.append(["Incident Analysis"])
    ws3.append([])
    ws3.append(["Metric", "Value"])
    _style_header(ws3, 3, 2)
    for k, v in data["incidents"].items():
        if isinstance(v, (int, float, str)) or v is None:
            ws3.append([k.replace("_", " ").title(), v])
        elif isinstance(v, dict):
            for kk, vv in v.items():
                ws3.append([f"  {kk}", vv])
    ws3.column_dimensions["A"].width = 28

    # ── High Risk Employees sheet ───────────────────────────
    ws4 = wb.create_sheet("High Risk Employees")
    ws4.append(["Name", "Department", "Designation", "Risk Score", "Risk Level"])
    _style_header(ws4, 1, 5)
    for e in data["high_risk_employees"]:
        ws4.append(
            [
                e["name"], e.get("department") or "", e.get("designation") or "",
                e["risk_score"], e["risk_level"],
            ]
        )
    for col, w in zip("ABCDE", (28, 18, 18, 12, 12)):
        ws4.column_dimensions[col].width = w

    # ── Threat Assessment sheet ─────────────────────────────
    ta = data.get("threat_assessment")
    if ta:
        ws5 = wb.create_sheet("Threat Assessment")
        ws5.append(["Employee", "Department", "Threat Score", "Threat Level"])
        _style_header(ws5, 1, 4)
        for t in ta["top_threats"]:
            ws5.append(
                [
                    t.get("employee_name") or "", t.get("department") or "",
                    t.get("threat_score", ""), t.get("threat_level") or "",
                ]
            )
        for col, w in zip("ABCD", (28, 18, 14, 14)):
            ws5.column_dimensions[col].width = w

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def employee_report_xlsx_bytes(
    db: Session, employee_id: str, days: int = 30
) -> bytes:
    """Render a single-employee report as an Excel workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    data = _employee_data(db, employee_id, days)
    emp = data["employee"]

    header_fill = PatternFill("solid", fgColor="1E293B")
    header_font = Font(color="FFFFFF", bold=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Employee"
    ws.append([f"ITBIS Employee Report — {emp['name']}"])
    ws.append(
        [
            f"{emp.get('department') or ''} — {emp.get('designation') or ''} "
            f"| Generated {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}"
        ]
    )
    ws.append([])
    ws.append(["Metric", "Value"])
    for cell in ("A4", "B4"):
        ws[cell].fill = header_fill
        ws[cell].font = header_font
    ws.append(["Total Alerts", data["total_alerts"]])
    ws.append(["Total Activity Logs", data["total_activity_logs"]])
    if data.get("latest_risk_score"):
        lr = data["latest_risk_score"]
        ws.append(["Latest Risk Score", f"{lr['score']} ({lr['level']})"])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 34

    ws2 = wb.create_sheet("Alert Breakdown")
    ws2.append(["Severity", "Count"])
    for cell in ("A1", "B1"):
        ws2[cell].fill = header_fill
        ws2[cell].font = header_font
    for k, v in (data.get("alert_severity_breakdown") or {}).items():
        ws2.append([k, v])
    # "Anomaly Type" header lands after severity header + data rows + blank line
    anomaly_header_row = 1 + len(data.get("alert_severity_breakdown") or {}) + 2
    ws2.append([])
    ws2.append(["Anomaly Type", "Count"])
    for cell in (f"A{anomaly_header_row}", f"B{anomaly_header_row}"):
        ws2[cell].fill = header_fill
        ws2[cell].font = header_font
    for k, v in (data.get("alert_type_breakdown") or {}).items():
        ws2.append([k, v])
    ws2.column_dimensions["A"].width = 22

    ta = data.get("threat_assessment") or {}
    if ta:
        ws3 = wb.create_sheet("Threat Models")
        ws3.append(["Model", "Score", "Level"])
        for cell in ("A1", "B1", "C1"):
            ws3[cell].fill = header_fill
            ws3[cell].font = header_font
        for name, m in (ta.get("model_scores") or {}).items():
            ws3.append([name.replace("_", " "), m.get("score", ""), m.get("level", "")])
        for col, w in zip("ABC", (24, 12, 12)):
            ws3.column_dimensions[col].width = w

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ── Filename helpers ────────────────────────────────────────────


def anomaly_filename(ext: str, days: int) -> str:
    return f"itbis_anomaly_report_{days}d_{datetime.utcnow():%Y%m%d}.{ext}"


def employee_filename(ext: str) -> str:
    return f"itbis_employee_report_{datetime.utcnow():%Y%m%d}.{ext}"
